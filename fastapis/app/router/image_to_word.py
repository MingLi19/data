import io
import os

from docx import Document
from docx.shared import Inches
from fastapi import APIRouter, File, Form, UploadFile
from PIL import Image

api = APIRouter()


@api.post("/image_to_word", summary="将图片转换为Word文档")
async def image_to_word(
    file: UploadFile = File(..., description="上传的图片文件"),
    description: str = Form(..., description="图片的文字描述"),
    save_path: str = Form("output.docx", description="文档保存路径"),
):
    # 创建Word文档
    document = Document()

    # 读取图片
    image = Image.open(io.BytesIO(await file.read()))

    # 将图片保存到临时文件
    temp_image_path = "temp_image.png"
    image.save(temp_image_path)

    # 将图片添加到Word文档
    document.add_picture(temp_image_path, width=Inches(6))

    # 添加自定义的文字描述
    if description:
        document.add_paragraph(description)

    # 删除临时文件
    os.remove(temp_image_path)

    # 保存Word文档到指定路径
    document.save(save_path)

    # 返回保存成功的消息
    return {"message": f"Document saved to {save_path}"}
